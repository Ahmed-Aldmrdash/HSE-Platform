#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_knowledge_base.py
========================
يبني ملف knowledge_base.json من مجلد refdata/ (SDS + drills + b1/b2 monthly
inspections) ليستخدمه lib/chatbot-kb.js في الشات بوت (الطبقة الأولى: Q&A
مقيّد بمصادر رسمية). كل مُدخل (entry) يحمل مصدره الأصلي (اسم الملف) حتى يقدر
البوت يستشهد بمصدر إجابته بدل ما يخترع من عنده.

يشتغل مرة واحدة الآن (بناء أولي) — أي تحديث لاحق لملفات SDS/الفحص الشهري
يحتاج إعادة تشغيل هذا السكريبت يدويًا، أو ربطه بزر "تحديث قاعدة المعرفة" في
لوحة الأدمن مستقبلاً (فكرة موجودة في تقرير التسليم).
"""
import json
import re
import sys
from pathlib import Path

import openpyxl
import docx

REFDATA = Path("/home/claude/refdata")
OUT_PATH = Path("/home/claude/work/platform/data/knowledge_base.json")

entries = []
entry_id = 0


def next_id(prefix):
    global entry_id
    entry_id += 1
    return f"{prefix}-{entry_id:04d}"


def decode_hashu(s):
    """يفك ترميز أسماء الملفات العربية اللي طلعت من unzip بصيغة '#U0627#U0644...'
    بدل الحروف العربية الأصلية (مشكلة ترميز في أداة فك الضغط المستخدمة)."""
    return re.sub(r"#U([0-9A-Fa-f]{4})", lambda m: chr(int(m.group(1), 16)), str(s))


def clean(s):
    if s is None:
        return ""
    return re.sub(r"\s+", " ", decode_hashu(s)).strip()


def sheet_to_lines(ws, max_rows=400):
    lines = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, max_rows), values_only=True):
        vals = [clean(c) for c in row if c is not None and clean(c) != ""]
        if vals:
            lines.append(" | ".join(vals))
    return lines


# ---------------------------------------------------------------------------
# 1) SDS — قوائم بيان المواد الكيميائية (Safety Data Sheets)
# ---------------------------------------------------------------------------
sds_dir = REFDATA / "SDS" / "SDS"
sds_count = 0
for f in sorted(sds_dir.glob("*.xlsx")) + sorted(sds_dir.glob("*.XLSX")):
    if f.name.startswith("drill_attendance"):
        continue  # ملف حضور تدريب اتحط بالغلط في مجلد SDS، مش SDS فعلي
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.worksheets[0]
        lines = sheet_to_lines(ws)
        if not lines:
            continue
        # اسم المادة: أول قيمة فعلية بعد سطر العنوان، أو اسم الملف نفسه
        material_name = clean(f.stem)
        for ln in lines:
            if "الاســـم التجــــــــــارى" in ln or "الاسم التجاري" in ln:
                parts = ln.split(":")
                if len(parts) > 1 and parts[-1].strip():
                    material_name = parts[-1].strip()
                break
        text = "\n".join(lines)
        entries.append({
            "id": next_id("SDS"),
            "type": "sds",
            "title": f"بيانات السلامة الكيميائية (SDS): {material_name}",
            "source_file": f"SDS/{clean(f.name)}",
            "tags": ["sds", "كيماوي", "مادة", material_name.lower()],
            "content": text,
        })
        sds_count += 1
    except Exception as e:
        print(f"[WARN] SDS parse failed: {f.name}: {e}", file=sys.stderr)

# ---------------------------------------------------------------------------
# 2) تجارب الطوارئ — تقارير تجارب إخلاء/حريق/إسعافات (drill docx reports)
# ---------------------------------------------------------------------------
drills_dir = REFDATA / "drills"
drill_count = 0
for f in sorted(drills_dir.rglob("*.docx")):
    try:
        d = docx.Document(f)
        paras = [clean(p.text) for p in d.paragraphs if clean(p.text)]
        for t in d.tables:
            for row in t.rows:
                cells = [clean(c.text) for c in row.cells if clean(c.text)]
                if cells:
                    paras.append(" | ".join(cells))
        if not paras:
            continue
        title = clean(f.stem)
        parent = clean(f.parent.name)
        text = "\n".join(paras)
        entries.append({
            "id": next_id("DRILL"),
            "type": "drill_report",
            "title": f"تقرير تجربة طوارئ: {title}",
            "source_file": f"drills/{parent}/{clean(f.name)}",
            "tags": ["تجربة طوارئ", "اخلاء", "طوارئ", parent.lower()],
            "content": text,
        })
        drill_count += 1
    except Exception as e:
        print(f"[WARN] drill parse failed: {f.name}: {e}", file=sys.stderr)

# ---------------------------------------------------------------------------
# 3) الفحص الشهري (b1 / b2) — قوائم معدات السلامة وحالتها بالفحص الأخير
# ---------------------------------------------------------------------------
def process_monthly(dir_path, area_label):
    count = 0
    for f in sorted(dir_path.rglob("*.xlsx")) + sorted(dir_path.rglob("*.XLSX")):
        try:
            wb = openpyxl.load_workbook(f, data_only=True)
            sheet_names = wb.sheetnames
            # آخر شيت غالبًا أحدث شهر مسجّل؛ لو فيه شيت اسمه 'list' نضيفه كمرجع
            base_lines = []
            if "list" in sheet_names:
                base_lines = sheet_to_lines(wb["list"], max_rows=250)
            latest_ws = wb[sheet_names[-1]]
            latest_lines = sheet_to_lines(latest_ws, max_rows=250)
            equip_name = clean(f.stem)
            text_parts = []
            if base_lines:
                text_parts.append("قائمة المعدات/المواقع المسجّلة:\n" + "\n".join(base_lines[:120]))
            if latest_lines:
                text_parts.append(f"آخر فحص مسجّل (شيت: {sheet_names[-1]}):\n" + "\n".join(latest_lines[:150]))
            if not text_parts:
                continue
            entries.append({
                "id": next_id("INSP"),
                "type": "monthly_inspection",
                "title": f"الفحص الشهري ({area_label}) — {equip_name}",
                "source_file": f"{dir_path.name}/{clean(f.relative_to(dir_path))}",
                "tags": ["فحص شهري", "تفتيش", area_label.lower(), equip_name.lower()],
                "content": "\n\n".join(text_parts),
            })
            count += 1
        except Exception as e:
            print(f"[WARN] inspection parse failed: {f.name}: {e}", file=sys.stderr)
    return count

b1_count = process_monthly(REFDATA / "b1_monthly", "b1")
b2_count = process_monthly(REFDATA / "b2_monthly", "b2")

OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
with open(OUT_PATH, "w", encoding="utf-8") as fh:
    json.dump({
        "generatedAt": "2026-09-12",
        "counts": {"sds": sds_count, "drills": drill_count, "b1_inspections": b1_count, "b2_inspections": b2_count},
        "entries": entries,
    }, fh, ensure_ascii=False, indent=1)

print(f"SDS: {sds_count} | drills: {drill_count} | b1: {b1_count} | b2: {b2_count} | total entries: {len(entries)}")
print(f"Output: {OUT_PATH} ({OUT_PATH.stat().st_size/1024:.1f} KB)")
